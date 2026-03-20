"""Unit tests for excel exporter module."""

import pytest
import pandas as pd
import tempfile
import os
from openpyxl import load_workbook

from src.excel_exporter import make_wa_url, export_to_excel, format_date


class TestFormatDate:
    """Tests for format_date function."""

    def test_format_iso_date(self):
        result = format_date("2024-01-15")
        assert result == "2024-01-15"

    def test_format_date_with_timezone(self):
        result = format_date("2024-01-15T00:00:00.000Z")
        assert result == "2024-01-15"

    def test_format_date_with_time(self):
        result = format_date("2024-01-15T18:30:30.721Z")
        assert result == "2024-01-15"

    def test_format_different_month(self):
        result = format_date("2024-12-25T10:00:00.000Z")
        assert result == "2024-12-25"

    def test_format_empty_string(self):
        result = format_date("")
        assert result == ""

    def test_format_nan_string(self):
        result = format_date("nan")
        assert result == ""

    def test_format_invalid_date_returns_original(self):
        result = format_date("invalid-date")
        assert result == "invalid-date"

    def test_format_none(self):
        result = format_date(None)
        assert result == ""


class TestMakeWaUrl:
    """Tests for make_wa_url function."""

    def test_local_format_phone(self):
        result = make_wa_url("01012345678", "Hello")
        assert result == "https://wa.me/201012345678?text=Hello"

    def test_international_format_phone(self):
        result = make_wa_url("+201012345678", "Hello")
        assert result == "https://wa.me/201012345678?text=Hello"

    def test_phone_with_spaces(self):
        result = make_wa_url("010 1234 5678", "Hello")
        assert result == "https://wa.me/201012345678?text=Hello"

    def test_international_with_country_code(self):
        result = make_wa_url("201012345678", "Hello")
        assert result == "https://wa.me/201012345678?text=Hello"

    def test_invalid_phone_returns_empty(self):
        result = make_wa_url("01012345678?", "Hello")
        assert result == ""

    def test_encodes_arabic_message(self):
        result = make_wa_url("01012345678", "مرحبا")
        assert "%D9%85%D8%B1%D8%AD%D8%A8%D8%A7" in result

    def test_message_with_special_characters(self):
        result = make_wa_url("01012345678", "Hello World!")
        assert "Hello%20World%21" in result


class TestExportToExcel:
    """Tests for export_to_excel function."""

    def test_creates_excel_file(self):
        df = pd.DataFrame(
            [
                {
                    "PhoneNumber": "01012345678",
                    "Message": "Test message",
                    "Date": "2024-01-15",
                }
            ]
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            result = export_to_excel(df, tmp_path)
            assert result == tmp_path
            assert os.path.exists(tmp_path)
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def test_excel_contains_headers(self):
        df = pd.DataFrame(
            [
                {
                    "PhoneNumber": "01012345678",
                    "Message": "Test message",
                    "Date": "2024-01-15",
                }
            ]
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            export_to_excel(df, tmp_path)
            wb = load_workbook(tmp_path)
            ws = wb.active

            assert ws.cell(row=1, column=1).value == "PhoneNumber"
            assert ws.cell(row=1, column=2).value == "Message"
            assert ws.cell(row=1, column=3).value == "Date"
            assert ws.cell(row=1, column=4).value == "WhatsApp Link"
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def test_excel_contains_data(self):
        df = pd.DataFrame(
            [
                {
                    "PhoneNumber": "01012345678",
                    "Message": "Test message",
                    "Date": "2024-01-15",
                }
            ]
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            export_to_excel(df, tmp_path)
            wb = load_workbook(tmp_path)
            ws = wb.active

            assert ws.cell(row=2, column=1).value == "01012345678"
            assert ws.cell(row=2, column=2).value == "Test message"
            assert ws.cell(row=2, column=3).value == "2024-01-15"
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def test_invalid_phone_shows_invalid_label(self):
        df = pd.DataFrame(
            [
                {
                    "PhoneNumber": "01012345678?",
                    "Message": "Test message",
                    "Date": "2024-01-15",
                }
            ]
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            export_to_excel(df, tmp_path)
            wb = load_workbook(tmp_path)
            ws = wb.active

            assert ws.cell(row=2, column=4).value == "Invalid phone"
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def test_valid_phone_creates_hyperlink(self):
        df = pd.DataFrame(
            [
                {
                    "PhoneNumber": "01012345678",
                    "Message": "Test message",
                    "Date": "2024-01-15",
                }
            ]
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            export_to_excel(df, tmp_path)
            wb = load_workbook(tmp_path)
            ws = wb.active

            cell = ws.cell(row=2, column=4)
            assert cell.value == "Send"
            assert cell.hyperlink is not None
            assert "wa.me" in cell.hyperlink.target
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def test_multiple_rows(self):
        df = pd.DataFrame(
            [
                {
                    "PhoneNumber": "01012345678",
                    "Message": "Message 1",
                    "Date": "2024-01-15",
                },
                {
                    "PhoneNumber": "01098765432",
                    "Message": "Message 2",
                    "Date": "2024-01-16",
                },
                {
                    "PhoneNumber": "01011111111",
                    "Message": "Message 3",
                    "Date": "2024-01-17",
                },
            ]
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            export_to_excel(df, tmp_path)
            wb = load_workbook(tmp_path)
            ws = wb.active

            assert ws.cell(row=2, column=1).value == "01012345678"
            assert ws.cell(row=3, column=1).value == "01098765432"
            assert ws.cell(row=4, column=1).value == "01011111111"
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def test_empty_dataframe(self):
        df = pd.DataFrame(columns=["PhoneNumber", "Message", "Date"])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            export_to_excel(df, tmp_path)
            wb = load_workbook(tmp_path)
            ws = wb.active

            assert ws.cell(row=1, column=1).value == "PhoneNumber"
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)
